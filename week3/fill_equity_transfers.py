#!/usr/bin/env python3
"""
补足股权转让数据
=================
从 manual_gold Excel 的 "4_股权转让" Sheet 中提取转让事件，
转换为 equity_transfer JSONL 记录，补充到 data/ 目录下。

原因: 当前 8 家公司的 JSONL 中均无 equity_transfer 记录，
      但赛分科技等公司的 Gold 中已有人工标注的转让数据。

输出: data/{code}_{name}_transfers.jsonl (独立传输文件)
      data/{code}_{name}_full.jsonl (含增资+快照+转让的完整文件)
"""

import json
import os
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

# ── Configuration ──
SCRIPT_DIR = Path(__file__).resolve().parent
WEEK3_DIR = SCRIPT_DIR  # week3/
DATA_DIR = WEEK3_DIR / "data"
GOLD_DIR = WEEK3_DIR / "manual_gold"
LOG_DIR = WEEK3_DIR / "outputs" / "logs"

os.makedirs(LOG_DIR, exist_ok=True)

COMPANIES = [
    ("001282", "三联锻造"),
    ("301563", "云汉芯城"),
    ("301581", "黄山谷捷"),
    ("603418", "友升股份"),
    ("688758", "赛分科技"),
    ("688775", "影石创新"),
    ("920100", "三协电机"),
    ("920116", "星图测控"),
]

# ── 列名标准化映射 ──
# Gold 各公司 Sheet 4 的列名可能略有不同，这里做映射
COLUMN_VARIANTS = {
    "pdf_page": ["PDF页码", "页码"],
    "transfer_date": ["转让日期", "日期"],
    "batch_label": ["批次标签", "批次"],
    "transferor": ["转让方", "出让人"],
    "transferee": ["受让方", "受让人"],
    "transfer_shares": ["转让注册资本(万元)", "转让股数(万股)", "转让数量(万股)"],
    "transfer_price_per_share": ["转让单价(元/注册资本)", "每股价格(元/股)"],
    "transfer_total_amount": ["转让总价(万元)", "总金额(万元)"],
    "transfer_reason": ["转让原因", "转让类型"],
    "source_label": ["来源标注"],
    "evidence_text": ["原文证据"],
    "transfer_type": ["转让类型"],
}


def find_column(col_name, headers):
    """按变体名匹配列索引。"""
    variants = COLUMN_VARIANTS.get(col_name, [col_name])
    for v in variants:
        if v in headers:
            return headers.index(v)
    return None


def extract_transfers(code, name):
    """从 Gold Excel 提取股权转让记录。"""
    gold_path = GOLD_DIR / f"{code}_{name}_gold.xlsx"
    if not gold_path.exists():
        return []

    wb = openpyxl.load_workbook(str(gold_path), data_only=True)
    sheet_names = wb.sheetnames

    # 查找转让 sheet (可能是 "4_股权转让" 或类似)
    transfer_sheet = None
    for sn in sheet_names:
        if "转让" in sn or "transfer" in sn.lower():
            transfer_sheet = sn
            break

    if not transfer_sheet:
        return []

    ws = wb[transfer_sheet]
    if ws.max_row <= 1:
        return []

    # 读表头
    headers = [str(ws.cell(1, c).value) if ws.cell(1, c).value else "" for c in range(1, ws.max_column + 1)]

    # 找各列索引
    idx_pdf_page = find_column("pdf_page", headers)
    idx_date = find_column("transfer_date", headers)
    idx_batch = find_column("batch_label", headers)
    idx_transferor = find_column("transferor", headers)
    idx_transferee = find_column("transferee", headers)
    idx_shares = find_column("transfer_shares", headers)
    idx_price = find_column("transfer_price_per_share", headers)
    idx_amount = find_column("transfer_total_amount", headers)
    idx_reason = find_column("transfer_reason", headers)
    idx_source = find_column("source_label", headers)
    idx_evidence = find_column("evidence_text", headers)
    idx_type = find_column("transfer_type", headers)

    records = []
    for row_idx in range(2, ws.max_row + 1):
        transferor = str(ws.cell(row_idx, idx_transferor + 1).value).strip() if idx_transferor is not None else ""
        transferee = str(ws.cell(row_idx, idx_transferee + 1).value).strip() if idx_transferee is not None else ""

        # 跳过空行/无效行
        if not transferor or transferor in ("无", "经逐页核查", "翻遍了"):
            continue
        if transferor.startswith("翻遍") or transferor.startswith("经逐页"):
            continue

        def safe_val(idx):
            if idx is None:
                return None
            v = ws.cell(row_idx, idx + 1).value
            if v is None:
                return None
            try:
                return float(v)
            except (ValueError, TypeError):
                return None

        rec = {
            "record_type": "equity_transfer",
            "stock_code": code,
            "company_name": name,
            "pdf_page": str(ws.cell(row_idx, idx_pdf_page + 1).value).strip() if idx_pdf_page is not None else "",
            "transfer_date": str(ws.cell(row_idx, idx_date + 1).value).strip() if idx_date is not None else "",
            "batch_label": str(ws.cell(row_idx, idx_batch + 1).value).strip() if idx_batch is not None else "",
            "transferor_name": transferor,
            "transferee_name": transferee,
            "transfer_shares_wan": safe_val(idx_shares),
            "transfer_price_per_share": safe_val(idx_price),
            "transfer_total_amount": safe_val(idx_amount),
            "transfer_reason": str(ws.cell(row_idx, idx_reason + 1).value).strip() if idx_reason is not None else "",
            "source_label": str(ws.cell(row_idx, idx_source + 1).value).strip() if idx_source is not None else "",
            "evidence_text": str(ws.cell(row_idx, idx_evidence + 1).value).strip() if idx_evidence is not None else "",
            "transfer_type": str(ws.cell(row_idx, idx_type + 1).value).strip() if idx_type is not None else "",
            "data_source": "manual_gold",
            "notes": "",
        }

        # 过滤: 如果转让方是"无/核查"类文本则跳过
        if any(skip in transferor for skip in ["无", "翻遍", "核查"]):
            continue

        records.append(rec)

    return records


def main():
    print("=" * 60)
    print("补足股权转让数据 - 从 Gold Excel 提取")
    print(f"执行时间: {datetime.now().isoformat()}")
    print("=" * 60)

    all_transfers = []
    company_summary = []

    for code, name in COMPANIES:
        transfers = extract_transfers(code, name)
        all_transfers.extend(transfers)
        company_summary.append((code, name, len(transfers)))

        # 写独立 JSONL
        if transfers:
            out_path = DATA_DIR / f"{code}_{name}_transfers.jsonl"
            with open(out_path, "w", encoding="utf-8") as f:
                for rec in transfers:
                    f.write(json.dumps(rec, ensure_ascii=False) + "\n")
            print(f"  {code} {name}: {len(transfers)} 条转让事件 -> {out_path.name}")
        else:
            print(f"  {code} {name}: 0 条转让事件")

    print(f"\n总计: {len(all_transfers)} 条转让事件")

    # ── 合并到完整 JSONL ──
    for code, name in COMPANIES:
        src_path = DATA_DIR / f"{code}_{name}.jsonl"
        if not src_path.exists():
            continue

        # 读原有记录
        existing = []
        with open(src_path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line:
                    existing.append(json.loads(line))

        # 追加转让记录
        transfers = [r for r in all_transfers if r["stock_code"] == code]
        full_records = existing + transfers

        full_path = DATA_DIR / f"{code}_{name}_full.jsonl"
        with open(full_path, "w", encoding="utf-8") as f:
            for rec in full_records:
                f.write(json.dumps(rec, ensure_ascii=False) + "\n")
        print(f"  [合并] {code} {name}: {len(existing)} + {len(transfers)} = {len(full_records)} -> {full_path.name}")

    # ── 汇总 ──
    print(f"\n各公司转让统计:")
    for code, name, count in company_summary:
        print(f"  {code} {name}: {count} 条")

    # 写日志
    log_path = LOG_DIR / "fill_equity_transfers.log"
    with open(log_path, "w", encoding="utf-8") as f:
        f.write(f"补足股权转让数据日志\n")
        f.write(f"时间: {datetime.now().isoformat()}\n\n")
        for code, name, count in company_summary:
            f.write(f"{code} {name}: {count} 条\n")
    print(f"\n日志: {log_path}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
