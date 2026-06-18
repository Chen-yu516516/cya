#!/usr/bin/env python3
"""
Week 3 Pipeline — 自动化提取与校验
=====================================
功能:
  1. 从 week3/data/ 读取 8 家公司的 JSONL 数据
  2. Schema 校验（类型检查 / 必填字段 / 数值合理性）
  3. 输出校验日志到 week3/outputs/logs/validation.log
  4. 转换 JSONL 为双 Sheet Excel，保存到 week3/outputs/auto_excel/
  5. 输出未经手工修正的 auto_jsonl 到 week3/outputs/auto_jsonl/

重要: 赛分科技 (688758) 在 auto 输出中保持原始错误（股而非万股），
      以展示 pipeline 在不依赖手工修正时的局限性。
"""

import json
import os
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("openpyxl not found, installing...")
    os.system(f"{sys.executable} -m pip install openpyxl -q")
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# ── Path configuration ──
SCRIPT_DIR = Path(__file__).resolve().parent
WEEK3_DIR = SCRIPT_DIR.parent  # week3/
DATA_DIR = WEEK3_DIR / "data"
OUTPUTS_DIR = WEEK3_DIR / "outputs"
LOG_DIR = OUTPUTS_DIR / "logs"
AUTO_EXCEL_DIR = OUTPUTS_DIR / "auto_excel"
AUTO_JSONL_DIR = OUTPUTS_DIR / "auto_jsonl"

os.makedirs(LOG_DIR, exist_ok=True)
os.makedirs(AUTO_EXCEL_DIR, exist_ok=True)
os.makedirs(AUTO_JSONL_DIR, exist_ok=True)

# ── Company registry ──
COMPANIES = [
    {"code": "001282", "name": "三联锻造"},
    {"code": "301563", "name": "云汉芯城"},
    {"code": "301581", "name": "黄山谷捷"},
    {"code": "603418", "name": "友升股份"},
    {"code": "688758", "name": "赛分科技"},
    {"code": "688775", "name": "影石创新"},
    {"code": "920100", "name": "三协电机"},
    {"code": "920116", "name": "星图测控"},
]

# ── Schema definitions ──
SUBSCRIPTION_FLOW_REQUIRED = [
    "record_type", "stock_code", "company_name", "pdf_page",
    "subscription_date", "subscriber_name", "subscription_amount_wan"
]

SUBSCRIPTION_FLOW_NUMERIC = [
    "subscription_shares_wan", "subscription_amount_wan", "subscription_price_yuan"
]

EQUITY_SNAPSHOT_REQUIRED = [
    "record_type", "stock_code", "company_name", "pdf_page",
    "time_point", "shareholder_name"
]

EQUITY_SNAPSHOT_NUMERIC = [
    "total_shares_wan", "total_capital_wan",
    "shares_held_wan", "capital_contribution_wan", "shareholding_ratio"
]

# ── Log buffer ──
log_lines: list[str] = []


def log(msg: str) -> None:
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = f"[{ts}] {msg}"
    log_lines.append(line)
    print(line)


# ── Step 1: Load JSONL ──
def load_jsonl(filepath: Path) -> list[dict]:
    records = []
    with open(filepath, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError as e:
                log(f"  [ERROR] JSON 解析失败: {e} | line: {line[:80]}")
    return records


# ── Step 2: Schema validation ──
def validate_record(rec: dict, idx: int, company: str) -> list[str]:
    errors = []
    rtype = rec.get("record_type")

    # Common checks
    for field in ["record_type", "stock_code", "company_name", "pdf_page"]:
        if field not in rec or rec[field] is None:
            errors.append(f"[{company}#{idx}] 缺少必填字段: {field}")

    if rtype not in ("subscription_flow", "equity_snapshot"):
        errors.append(f"[{company}#{idx}] 未知 record_type: {rtype}")
        return errors

    if rtype == "subscription_flow":
        for field in SUBSCRIPTION_FLOW_REQUIRED:
            if field not in rec or rec[field] is None:
                errors.append(f"[{company}#{idx}] subscription_flow 缺少必填字段: {field}")
        # Numeric checks
        for field in SUBSCRIPTION_FLOW_NUMERIC:
            val = rec.get(field)
            if val is not None and not isinstance(val, (int, float)):
                errors.append(f"[{company}#{idx}] {field} 类型错误: 期望数值，实际 {type(val).__name__}")
        # Reasonability checks
        shares = rec.get("subscription_shares_wan")
        if shares is not None and isinstance(shares, (int, float)) and shares < 0:
            errors.append(f"[{company}#{idx}] subscription_shares_wan 为负数: {shares}")
        amount = rec.get("subscription_amount_wan")
        if amount is not None and isinstance(amount, (int, float)) and amount < 0:
            errors.append(f"[{company}#{idx}] subscription_amount_wan 为负数: {amount}")
        # Detect possible unit error: shares_wan > 100000 suggests raw 股 not 万股
        if shares is not None and isinstance(shares, (int, float)) and shares > 100000:
            errors.append(f"[{company}#{idx}] [WARN] subscription_shares_wan 值异常大 ({shares})，疑似单位为股而非万股")

    elif rtype == "equity_snapshot":
        for field in EQUITY_SNAPSHOT_REQUIRED:
            if field not in rec or rec[field] is None:
                errors.append(f"[{company}#{idx}] equity_snapshot 缺少必填字段: {field}")
        for field in EQUITY_SNAPSHOT_NUMERIC:
            val = rec.get(field)
            if val is not None and not isinstance(val, (int, float)):
                errors.append(f"[{company}#{idx}] {field} 类型错误: 期望数值，实际 {type(val).__name__}")
        # shareholding_ratio should be 0-100
        ratio = rec.get("shareholding_ratio")
        if ratio is not None and isinstance(ratio, (int, float)):
            if ratio < 0 or ratio > 100:
                errors.append(f"[{company}#{idx}] shareholding_ratio 超出合理范围 [0,100]: {ratio}")

    return errors


# ── Step 3: Excel generation ──
def write_excel(
    filepath: Path,
    company_name: str,
    stock_code: str,
    subscription_records: list[dict],
    equity_records: list[dict],
) -> None:
    wb = openpyxl.Workbook()

    # Styles
    header_font = Font(name="微软雅黑", bold=True, size=11)
    header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def style_header(ws, headers, col_widths):
        for col_idx, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w

    # ── Sheet 1: 认缴流量 ──
    ws1 = wb.active
    ws1.title = "1_认缴流量"
    flow_headers = [
        "PDF页码", "增资日期", "批次标签", "认购方",
        "认购数量(万股)", "认购金额(万元)", "认购价格(元/股)", "原文证据"
    ]
    flow_widths = [10, 16, 14, 24, 18, 18, 18, 80]
    style_header(ws1, flow_headers, flow_widths)

    for row_idx, rec in enumerate(subscription_records, 2):
        ws1.cell(row=row_idx, column=1, value=rec.get("pdf_page"))
        ws1.cell(row=row_idx, column=2, value=rec.get("subscription_date"))
        ws1.cell(row=row_idx, column=3, value=rec.get("batch_label"))
        ws1.cell(row=row_idx, column=4, value=rec.get("subscriber_name"))
        ws1.cell(row=row_idx, column=5, value=rec.get("subscription_shares_wan"))
        ws1.cell(row=row_idx, column=6, value=rec.get("subscription_amount_wan"))
        ws1.cell(row=row_idx, column=7, value=rec.get("subscription_price_yuan"))
        ws1.cell(row=row_idx, column=8, value=(rec.get("evidence_text") or "")[:32767])
        for c in range(1, 9):
            ws1.cell(row=row_idx, column=c).alignment = cell_alignment
            ws1.cell(row=row_idx, column=c).border = thin_border

    # ── Sheet 2: 股权结构存量 ──
    ws2 = wb.create_sheet("2_股权结构存量")
    eq_headers = [
        "PDF页码", "时点", "股权结构口径", "总股本(万股)", "总出资额(万元)",
        "股东名称", "持股数(万股)", "出资额(万元)", "持股比例", "原文证据"
    ]
    eq_widths = [10, 22, 18, 16, 16, 22, 16, 16, 12, 80]
    style_header(ws2, eq_headers, eq_widths)

    for row_idx, rec in enumerate(equity_records, 2):
        ws2.cell(row=row_idx, column=1, value=rec.get("pdf_page"))
        ws2.cell(row=row_idx, column=2, value=rec.get("time_point"))
        ws2.cell(row=row_idx, column=3, value=rec.get("equity_structure_scope"))
        ws2.cell(row=row_idx, column=4, value=rec.get("total_shares_wan"))
        ws2.cell(row=row_idx, column=5, value=rec.get("total_capital_wan"))
        ws2.cell(row=row_idx, column=6, value=rec.get("shareholder_name"))
        ws2.cell(row=row_idx, column=7, value=rec.get("shares_held_wan"))
        ws2.cell(row=row_idx, column=8, value=rec.get("capital_contribution_wan"))
        ws2.cell(row=row_idx, column=9, value=rec.get("shareholding_ratio"))
        ws2.cell(row=row_idx, column=10, value=(rec.get("evidence_text") or "")[:32767])
        for c in range(1, 11):
            ws2.cell(row=row_idx, column=c).alignment = cell_alignment
            ws2.cell(row=row_idx, column=c).border = thin_border

    wb.save(filepath)


# ── Step 4: Write auto JSONL (raw copy, NO manual corrections) ──
def write_auto_jsonl(records: list[dict], filepath: Path, source_info: str) -> None:
    with open(filepath, "w", encoding="utf-8") as f:
        for rec in records:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")


# ── Main pipeline ──
def main():
    log("=" * 60)
    log("Week 3 Pipeline — 自动化提取与校验开始")
    log(f"执行时间: {datetime.now().isoformat()}")
    log("=" * 60)
    log("")

    total_subscription = 0
    total_equity = 0
    total_errors = 0
    total_warnings = 0

    for comp in COMPANIES:
        code = comp["code"]
        name = comp["name"]
        jsonl_path = DATA_DIR / f"{code}_{name}.jsonl"

        log(f"处理: {code} {name}")
        log(f"  数据源: {jsonl_path}")

        if not jsonl_path.exists():
            log(f"  [ERROR] 文件不存在: {jsonl_path}")
            continue

        # Load
        records = load_jsonl(jsonl_path)
        log(f"  读取记录: {len(records)} 条")

        # Split
        subscription_records = [r for r in records if r.get("record_type") == "subscription_flow"]
        equity_records = [r for r in records if r.get("record_type") == "equity_snapshot"]
        log(f"  认缴流量: {len(subscription_records)} 条 | 股权存量: {len(equity_records)} 条")
        total_subscription += len(subscription_records)
        total_equity += len(equity_records)

        # Validate
        company_errors = 0
        company_warnings = 0
        for idx, rec in enumerate(records):
            errs = validate_record(rec, idx, f"{code} {name}")
            for e in errs:
                if "[WARN]" in e:
                    company_warnings += 1
                    log(f"  {e}")
                else:
                    company_errors += 1
                    log(f"  {e}")

        total_errors += company_errors
        total_warnings += company_warnings

        # Write auto Excel
        excel_path = AUTO_EXCEL_DIR / f"{code}_{name}_auto.xlsx"
        write_excel(excel_path, name, code, subscription_records, equity_records)
        log(f"  生成 Excel: {excel_path}")

        # Write auto JSONL (raw copy — NO manual corrections)
        # For 赛分科技, this preserves the original error (股 instead of 万股)
        auto_jsonl_path = AUTO_JSONL_DIR / f"{code}_{name}_auto.jsonl"
        write_auto_jsonl(records, auto_jsonl_path, f"Source: {jsonl_path}")
        log(f"  生成 auto_jsonl: {auto_jsonl_path} (共 {len(records)} 条)")

        log("")

    # Summary
    log("=" * 60)
    log("Pipeline 执行完毕 — 汇总")
    log(f"  总记录数: {total_subscription + total_equity}")
    log(f"  认缴流量: {total_subscription} | 股权存量: {total_equity}")
    log(f"  Schema 错误: {total_errors} | 警告: {total_warnings}")
    log(f"  生成 Excel: {len(COMPANIES)} 份 -> {AUTO_EXCEL_DIR}")
    log(f"  生成 auto_jsonl: {len(COMPANIES)} 份 -> {AUTO_JSONL_DIR}")
    log("=" * 60)

    # Write validation log
    log_path = LOG_DIR / "validation.log"
    with open(log_path, "w", encoding="utf-8") as f:
        f.write("\n".join(log_lines))
    print(f"\nValidation log written to: {log_path}")


if __name__ == "__main__":
    main()
